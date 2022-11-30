from openpyxl import Workbook


t0 = doc.tables[0]

workbook = Workbook()
sheet = workbook.active

for i in range(len(t0.rows)):
    list1 = []
    for j in range(len(t0.columns)):
        list1.append(t0.cell(i,j).text)
    sheet.append(list1)
workbook.save(filename = r"E:\2022 年城建环资工作\代表建议议案\test1.xlsx")

